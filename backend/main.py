"""FastAPI application for the local MVP offer workflow."""
from __future__ import annotations

import base64
import io
import json
import os
import re
import secrets as _secrets
from contextlib import asynccontextmanager
from datetime import date
from pathlib import Path
from typing import Any, Literal, Optional

from fastapi import FastAPI, File, HTTPException, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from starlette.middleware.base import BaseHTTPMiddleware
from starlette.requests import Request as StarletteRequest
from starlette.responses import Response as StarletteResponse

from .config import Config, get_config, get_user_profile
from .generator import build_offer_prompt, format_date_german, generate_offer, parse_ai_response
from .renderer import render_offer

_ROOT = Path(__file__).parent.parent
_FRONTEND_DIR = _ROOT / "frontend"
_OUTPUT_DIR = _ROOT / "output"
_PROMPTS_DIR = _OUTPUT_DIR / "prompts"
_GENERATED_DIR = _OUTPUT_DIR / "offers"


@asynccontextmanager
async def lifespan(_: FastAPI):
    get_config()
    get_user_profile()
    _PROMPTS_DIR.mkdir(parents=True, exist_ok=True)
    _GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    yield


class _BasicAuthMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request: StarletteRequest, call_next):
        password = os.environ.get("APP_PASSWORD", "").strip()
        if not password:
            return await call_next(request)
        if request.method == "OPTIONS" or request.url.path == "/healthz":
            return await call_next(request)

        auth = request.headers.get("Authorization", "")
        if auth.startswith("Basic "):
            try:
                decoded = base64.b64decode(auth[6:]).decode("utf-8")
                username, _, pwd = decoded.partition(":")
                exp_user = os.environ.get("APP_USERNAME", "accantec")
                if _secrets.compare_digest(username, exp_user) and _secrets.compare_digest(pwd, password):
                    return await call_next(request)
            except Exception:
                pass

        return StarletteResponse(
            status_code=401,
            headers={"WWW-Authenticate": 'Basic realm="Angebotserstellung"'},
            content="Anmeldung erforderlich.",
        )


app = FastAPI(
    title="accantec Angebotserstellung",
    description="Lokale MVP-App fuer die accantec Angebotserstellung",
    version="2.0.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://localhost:4173",
        "http://localhost:8000",
        "http://127.0.0.1:8000",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.add_middleware(_BasicAuthMiddleware)

if _FRONTEND_DIR.exists():
    app.mount("/static", StaticFiles(directory=str(_FRONTEND_DIR)), name="static")


class LineItem(BaseModel):
    description: str
    unit: str = "Tag(e)"
    rate: float
    days: float


class OfferFormPayload(BaseModel):
    offer_number: str = Field(min_length=3)
    generation_date: Optional[str] = None
    validity_date: str
    customer_company: str = Field(min_length=1)
    customer_street: str = Field(min_length=1)
    customer_postal: str = Field(min_length=1)
    customer_city: str = Field(min_length=1)
    customer_country: str = Field(min_length=1)
    customer_requester_name: str = Field(min_length=1)
    customer_requester_email: str = Field(min_length=1)
    customer_requester_role: str = ""
    customer_primary_contact_gender: str = ""
    customer_primary_contact_name: str = Field(min_length=1)
    customer_primary_contact_email: str = Field(min_length=1)
    customer_primary_contact_phone: str = ""
    project_name: str = Field(min_length=1)
    project_context: str = Field(min_length=1)
    pricing_mode: str = "time-and-materials"
    location_mode: str = "remote"
    line_items: list[LineItem] = []
    leistungsausschluesse: Optional[str] = None


class PromptRequest(OfferFormPayload):
    pass


class PromptResponse(BaseModel):
    prompt: str
    prompt_path: str


class RenderRequest(BaseModel):
    mode: Literal["markdown", "html", "pdf"] = "markdown"
    offer_number: str = Field(min_length=3)
    markdown: Optional[str] = None
    form_data: Optional[OfferFormPayload] = None
    ai_response: str = ""


class RenderMarkdownResponse(BaseModel):
    markdown: str
    offer_number: str


@app.get("/", response_class=HTMLResponse)
async def serve_index() -> HTMLResponse:
    index_path = _FRONTEND_DIR / "index.html"
    if index_path.exists():
        return HTMLResponse(content=index_path.read_text(encoding="utf-8"))
    return HTMLResponse("<h1>accantec Angebotserstellung</h1>")


@app.get("/healthz")
async def healthz() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/api/profile")
async def get_profile() -> dict:
    config = get_config()
    user_profile = get_user_profile()
    return {
        "profile": {
            "id": user_profile.id,
            "name": user_profile.name,
            "title": user_profile.title,
            "email": user_profile.email,
            "phone": user_profile.phone,
            "company_name": user_profile.company_name,
            "signature_name": user_profile.signature_name,
            "signature_title": user_profile.signature_title,
            "closing": user_profile.closing,
        },
        "company": {
            "name": config.company.name,
            "address_line1": config.company.address_line1,
            "address_line2": config.company.address_line2,
            "city": config.company.city,
            "website": config.company.website,
        },
        "defaults": {
            "validity_weeks": config.defaults.validity_weeks,
            "location_mode": config.defaults.location_mode,
            "currency": config.defaults.currency,
        },
        "rates": {
            "managing_consultant": config.rates.managing_consultant,
            "senior_consultant": config.rates.senior_consultant,
            "consultant": config.rates.consultant,
            "junior_consultant": config.rates.junior_consultant,
        },
    }


@app.post("/api/prompt", response_model=PromptResponse)
async def create_prompt(req: PromptRequest) -> PromptResponse:
    config = get_config()
    user_profile = get_user_profile()
    form_data = _build_form_data(req)
    prompt = build_offer_prompt(form_data, config, user_profile)
    _PROMPTS_DIR.mkdir(parents=True, exist_ok=True)
    prompt_path = _PROMPTS_DIR / f"{req.offer_number}.md"
    prompt_path.write_text(prompt, encoding="utf-8")
    return PromptResponse(prompt=prompt, prompt_path=str(prompt_path))


@app.post("/api/render")
async def render(req: RenderRequest):
    config = get_config()
    user_profile = get_user_profile()

    if req.mode in ("pdf", "html"):
        if not req.markdown:
            raise HTTPException(status_code=422, detail="Markdown fehlt fuer den Export.")
        try:
            if req.mode == "pdf":
                pdf_path = _offer_dir(req.offer_number) / f"Angebot_{req.offer_number}.pdf"
                pdf_bytes = render_offer(req.markdown, config, output_path=pdf_path)
                return Response(
                    content=pdf_bytes,
                    media_type="application/pdf",
                    headers={
                        "Content-Disposition": f'attachment; filename="Angebot_{req.offer_number}.pdf"'
                    },
                )
            else:
                from .renderer import render_markdown_to_html
                html_str = render_markdown_to_html(req.markdown, config)
                html_path = _offer_dir(req.offer_number) / f"Angebot_{req.offer_number}.html"
                html_path.write_text(html_str, encoding="utf-8")
                return Response(
                    content=html_str.encode("utf-8"),
                    media_type="text/html",
                    headers={
                        "Content-Disposition": f'attachment; filename="Angebot_{req.offer_number}.html"'
                    },
                )
        except Exception as exc:
            raise HTTPException(status_code=500, detail=f"Rendering fehlgeschlagen: {exc}") from exc

    if req.form_data is None:
        raise HTTPException(status_code=422, detail="Formulardaten fehlen fuer die Markdown-Erzeugung.")

    try:
        form_data = _build_form_data(req.form_data)
        ai_response = (req.ai_response or "").strip()
        if ai_response:
            try:
                ai_content = parse_ai_response(ai_response)
            except ValueError as exc:
                raise HTTPException(status_code=422, detail=str(exc)) from exc
        else:
            ai_content = None
        markdown = generate_offer(form_data, config, user_profile, ai_content=ai_content)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Markdown-Erzeugung fehlgeschlagen: {exc}") from exc

    md_path = _offer_dir(req.offer_number) / f"Angebot_{req.offer_number}.md"
    md_path.write_text(markdown, encoding="utf-8")
    return RenderMarkdownResponse(markdown=markdown, offer_number=req.offer_number)


def _offer_dir(offer_number: str) -> Path:
    offer_dir = _GENERATED_DIR / offer_number
    offer_dir.mkdir(parents=True, exist_ok=True)
    return offer_dir


def _build_form_data(req: OfferFormPayload) -> dict:
    generation_date = req.generation_date or date.today().isoformat()
    primary_contact_last_name = req.customer_primary_contact_name.split()[-1]
    from dateutil.parser import parse as parse_date  # type: ignore

    validity_date_german = format_date_german(parse_date(req.validity_date).date())
    return {
        "offer_number": req.offer_number.strip(),
        "generation_date": generation_date,
        "validity_date": req.validity_date,
        "validity_date_german": validity_date_german,
        "customer_company": req.customer_company.strip(),
        "customer_company_legal_name": req.customer_company.strip(),
        "customer_street": req.customer_street.strip(),
        "customer_postal": req.customer_postal.strip(),
        "customer_city": req.customer_city.strip(),
        "customer_country": req.customer_country.strip(),
        "customer_requester_name": req.customer_requester_name.strip(),
        "customer_requester_email": req.customer_requester_email.strip(),
        "customer_requester_role": req.customer_requester_role.strip(),
        "customer_primary_contact_gender": req.customer_primary_contact_gender.strip(),
        "customer_primary_contact_name": req.customer_primary_contact_name.strip(),
        "customer_primary_contact_email": req.customer_primary_contact_email.strip(),
        "customer_primary_contact_phone": req.customer_primary_contact_phone.strip(),
        "customer_primary_contact_last_name": primary_contact_last_name,
        "customer_contact_person": req.customer_primary_contact_name.strip(),
        "customer_contact_last_name": primary_contact_last_name,
        "customer_gender": req.customer_primary_contact_gender.strip(),
        "project_name": req.project_name.strip(),
        "project_context": req.project_context.strip(),
        "pricing_mode": req.pricing_mode,
        "location_mode": req.location_mode,
        "line_items": [item.model_dump() for item in req.line_items],
        "leistungsausschluesse": req.leistungsausschluesse,
    }


# ─── File text extraction ────────────────────────────────────────────────────

@app.post("/api/extract-text")
async def extract_text(file: UploadFile = File(...)):
    content = await file.read()
    filename = file.filename or ""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    try:
        if ext in ("txt", "md", "markdown", "csv"):
            text = content.decode("utf-8", errors="replace")
        elif ext == "pdf":
            from pypdf import PdfReader  # type: ignore
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext in ("xlsx", "xls"):
            import openpyxl  # type: ignore
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            rows: list[str] = []
            for sheet in wb.worksheets:
                rows.append(f"## {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    if any(c is not None for c in row):
                        rows.append("\t".join("" if c is None else str(c) for c in row))
            text = "\n".join(rows)
        elif ext in ("pptx", "ppt"):
            from pptx import Presentation  # type: ignore
            prs = Presentation(io.BytesIO(content))
            slides: list[str] = []
            for i, slide in enumerate(prs.slides, 1):
                parts = [f"## Folie {i}"]
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        parts.append(shape.text.strip())
                slides.append("\n".join(parts))
            text = "\n\n".join(slides)
        elif ext in ("docx", "doc"):
            from docx import Document  # type: ignore
            doc = Document(io.BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        else:
            text = content.decode("utf-8", errors="replace")
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Datei konnte nicht gelesen werden: {exc}") from exc

    return {"text": text[:60000], "filename": filename}


# ─── Chat endpoint ────────────────────────────────────────────────────────────

class ChatMessagePayload(BaseModel):
    role: Literal["user", "assistant"]
    content: str


class ChatDraftPayload(BaseModel):
    project_name: str = ""
    customer_company: str = ""
    customer_street: str = ""
    customer_postal: str = ""
    customer_city: str = ""
    customer_primary_contact_gender: str = ""
    customer_primary_contact_name: str = ""
    customer_primary_contact_phone: str = ""
    customer_primary_contact_email: str = ""
    customer_requester_name: str = ""
    customer_requester_role: str = ""
    customer_requester_email: str = ""
    location_mode: str = "remote"
    cover_paragraphs: list[str] = []
    ausgangssituation: str = ""
    einschaetzung_rows: list[dict[str, Any]] = []
    projektziele: list[str] = []
    leistungsbeschreibung: str = ""
    leistungsausschluesse: str = ""
    line_items: list[dict[str, Any]] = []


class ChatOfferRequest(BaseModel):
    messages: list[ChatMessagePayload]
    draft: ChatDraftPayload


class ChatOfferResponse(BaseModel):
    reply: str
    updates: dict[str, Any] = {}


def _chat_system_prompt(config: Config, draft: ChatDraftPayload) -> str:
    generation_policy = config.generation_policy.strip()

    state: list[str] = []
    if draft.project_name:
        state.append(f"- Projektname: {draft.project_name}")
    if draft.customer_company:
        state.append(f"- Auftraggeber: {draft.customer_company}")
    if draft.customer_primary_contact_name:
        state.append(f"- Ansprechpartner: {draft.customer_primary_contact_name}")
    if draft.ausgangssituation:
        state.append("- Ausgangssituation: vorhanden")
    if draft.leistungsbeschreibung:
        state.append("- Leistungsbeschreibung: vorhanden")
    if draft.line_items:
        state.append(f"- Preispositionen: {len(draft.line_items)} Position(en)")
    state_block = "\n".join(state) if state else "- (Entwurf ist noch leer)"

    return f"""Du bist ein professioneller Assistent für die Angebotserstellung bei der accantec information solutions GmbH. Du hilfst beim Erstellen und Überarbeiten von Beratungsangeboten auf Deutsch.

AUSGABEFORMAT – antworte IMMER mit validem JSON, ohne Markdown-Blöcke:
{{"reply": "Deine Antwort auf Deutsch (Fließtext)", "updates": {{...nur geänderte Felder...}}}}

FELDER in "updates" (alle optional, nur setzen wenn geändert):
- project_name (String)
- customer_company, customer_street, customer_postal, customer_city (Strings)
- customer_primary_contact_gender ("Herr" | "Frau" | ""), customer_primary_contact_name, customer_primary_contact_phone, customer_primary_contact_email (Strings)
- customer_requester_name, customer_requester_role, customer_requester_email (Strings)
- location_mode ("remote" | "hybrid" | "onsite")
- cover_paragraphs (Liste mit genau 3 Strings – Anschreiben-Absätze, du-Form: "vielen Dank für...")
- ausgangssituation (Text, Absätze mit \\n\\n getrennt)
- einschaetzung_rows (Liste von {{"problem": "...", "empfehlung": "..."}}, 3-5 Zeilen)
- projektziele (Liste von Strings, je mit starkem Infinitiv beginnend, 4-6 Stück)
- leistungsbeschreibung (Markdown: ### Phasenname\\n\\n- Tätigkeit 1\\n- Tätigkeit 2)
- leistungsausschluesse (String, Standard "./.")
- line_items (Liste von {{"description": "...", "unit": "Tag(e)", "rate": 1300.0, "days": 4.0}})

WORKFLOW:
1. Wenn Kerndaten fehlen (Auftraggeber, Projektname, Kontext): stelle max. 1-2 gezielte Rückfragen. Lass "updates" dabei leer.
2. Sobald du Auftraggeber + Projektname + Projektkontext kennst: generiere cover_paragraphs, ausgangssituation, einschaetzung_rows, projektziele und leistungsbeschreibung vollständig in einem Schritt.
3. Für gezielte Änderungswünsche ("Ausgangssituation kürzer", "Phase 4 ergänzen", "Preis anpassen"): ändere nur das betroffene Feld.
4. Bestätige kurz auf Deutsch was du gemacht hast (1-2 Sätze in "reply").

QUALITÄTSSTANDARDS FÜR INHALTLICHE FELDER:
{generation_policy}

AKTUELLER ENTWURF:
{state_block}"""


@app.post("/api/chat/messages", response_model=ChatOfferResponse)
async def chat_offer_messages(req: ChatOfferRequest) -> ChatOfferResponse:
    api_key = os.environ.get("ANTHROPIC_API_KEY", "").strip()
    if not api_key:
        raise HTTPException(
            status_code=503,
            detail="ANTHROPIC_API_KEY ist nicht gesetzt. Bitte in .env konfigurieren.",
        )

    try:
        import anthropic  # type: ignore
    except ImportError as exc:
        raise HTTPException(
            status_code=503,
            detail="anthropic-Paket fehlt. Bitte 'pip install anthropic' ausführen.",
        ) from exc

    config = get_config()
    system_prompt = _chat_system_prompt(config, req.draft)

    client = anthropic.AsyncAnthropic(api_key=api_key)
    response = await client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=4096,
        system=system_prompt,
        messages=[{"role": m.role, "content": m.content} for m in req.messages],
    )

    raw = response.content[0].text.strip()

    # Strip markdown code fences if the model added them despite instructions
    raw = re.sub(r"^```(?:json)?\s*", "", raw)
    raw = re.sub(r"\s*```$", "", raw)

    try:
        data = json.loads(raw)
        reply: str = data.get("reply", "Verstanden.")
        updates: dict[str, Any] = data.get("updates", {})
    except json.JSONDecodeError:
        reply = raw
        updates = {}

    return ChatOfferResponse(reply=reply, updates=updates)
